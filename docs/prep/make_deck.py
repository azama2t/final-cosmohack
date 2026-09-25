# АРХИВ подготовительного этапа: генератор прежней деки (‰, живые снимки). Не запускать отсюда — пути рассчитаны на scripts/.
# Материалы кейса генерирует scripts/make_deck_case.py.
r"""Build the pitch materials from reports/final_numbers.json (the only source of numbers):

  reports/deck.pptx        10 slides, 16:9, dark theme of the UI, coral accent; speaker notes = speech of the slide
  docs/SPEECH.md           4-minute speech word for word, timing per slide, numbers to memorise, hard questions
  reports/qa.md            31 jury questions with short honest answers
  reports/deck_preview/    PNG per slide (with --preview, LibreOffice)

Usage (from repo root):
    .venv\Scripts\python.exe scripts\make_deck.py              # deck + SPEECH.md + qa.md
    .venv\Scripts\python.exe scripts\make_deck.py --preview    # + PNG previews

Numbers are never typed by hand: every figure below is a placeholder filled from final_numbers.json
(missing -> "—"). When final_numbers.json is regenerated (new main model, test computed), rerun this script.
Screenshots: reports/screens/final (L39 fresh set, scripts/screenshots.py --extra --l20), else docs/img; crops are
made in memory. Demo region (speech, DEMO table) = manifest.json `demo` {region, date} when present (L43,
build_service_data.py --demo-region; the tour opens the same region), else best_region of
reports/screens/final/result.json (same bestRegion rule as the tour), fallback: the rule recomputed from
service/data/manifest.json. Screenshot captions always name the region that is on the screenshots (result.json).
If the demo date has no drift, the speech switches to the drift region of the screenshots for the drift step.
Slide rules: the title states the conclusion; one big number; <= 25 words of body; captions on images; >= 18 pt body.
"""
from __future__ import annotations

import argparse
import io
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "reports" / "screens" / "final"
SCREEN_DIRS = [SHOTS]
# README copies of the same frames (docs/img, JPG): fallback when reports/screens/final is absent (clean clone)
IMG_FALLBACK = {"01_overview_1920.png": "overview.jpg", "02_region.png": "region.jpg",
                "03_detection_card.png": "detection_card.jpg", "09_drift.png": "drift.jpg",
                "21_zone_card.png": "zone_card.jpg"}

BG = RGBColor(0x0A, 0x11, 0x1F)
PANEL = RGBColor(0x12, 0x1C, 0x2E)
LINE = RGBColor(0x2A, 0x3A, 0x55)
TEXT = RGBColor(0xE8, 0xEE, 0xF6)
MUTED = RGBColor(0x9A, 0xAA, 0xC0)
CORAL = RGBColor(0xFF, 0x6B, 0x4A)
FONT = "Segoe UI"
MONO = "Consolas"
DASH = "—"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)


# ------------------------------------------------------------------ numbers
def g(d, path, default=None):
    cur = d
    for p in path.split("."):
        if isinstance(cur, dict) and p in cur:
            cur = cur[p]
        elif isinstance(cur, list):
            try:
                cur = cur[int(p)]
            except (ValueError, IndexError):
                return default
        else:
            return default
    return default if cur is None else cur


def num(v, nd=3):
    """Russian decimal comma; thin grouping for integers."""
    if v is None:
        return DASH
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    if nd == 0:
        return f"{int(round(x)):,}".replace(",", "\u00a0")
    return f"{x:.{nd}f}".replace(".", ",")


def plural(n, one, few, many):
    try:
        n = abs(int(n))
    except (TypeError, ValueError):
        return many
    if n % 10 == 1 and n % 100 != 11:
        return one
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return few
    return many


def pct(v):
    return DASH if v is None else f"{100 * float(v):.0f}\u00a0%"


def _flagged(d: dict | None) -> bool:
    q = (d or {}).get("quality") or {}
    return bool(q.get("haze") or q.get("glint_or_haze"))


def best_region_manifest() -> tuple[str | None, str | None]:
    """Python copy of bestRegion() (service/frontend/src/lib/data.ts): latest scene reliable (no haze/glint,
    clouds <= 50 %) and with detections, prefer drift, then most detections / drift / index.
    Returns (best id, id for the drift frame or None)."""
    mp = ROOT / "service" / "data" / "manifest.json"
    if not mp.exists():
        return None, None
    regs = json.loads(mp.read_text(encoding="utf-8")).get("regions") or []

    def summ(r):
        return r.get("summary") or {}

    def sdate(r):
        ld = summ(r).get("latest_date")
        return next((d for d in r["dates"] if d.get("date") == ld), r["dates"][-1] if r["dates"] else None)

    def det(r):
        return summ(r).get("n_detections") or 0

    def idx(r):
        v = summ(r).get("index_permille")
        return -1 if v is None else v

    def drift(r):
        return 1 if (sdate(r) or {}).get("drift") else 0

    clean = [r for r in regs if r.get("dates") and sdate(r) is not None and not _flagged(sdate(r))
             and ((sdate(r) or {}).get("cloud_frac") or 0) <= 0.5 and not summ(r).get("haze")
             and sdate(r) is r["dates"][-1] and det(r) > 0]
    pool = [r for r in clean if drift(r)] or clean
    if pool:
        best = sorted(pool, key=lambda r: (-det(r), -drift(r), -idx(r)))[0]
    else:
        best = sorted(regs, key=lambda r: (-det(r), -idx(r)))[0] if regs else None
    if not best:
        return None, None
    return best.get("id"), (best.get("id") if drift(best) else None)


def manifest_demo() -> tuple[dict | None, dict]:
    """L43: (manifest `demo` or None, {region id: {date: has drift}}) from service/data/manifest.json."""
    mp = ROOT / "service" / "data" / "manifest.json"
    if not mp.exists():
        return None, {}
    man = json.loads(mp.read_text(encoding="utf-8"))
    drift = {r.get("id"): {d.get("date"): bool(d.get("drift")) for d in r.get("dates", [])}
             for r in man.get("regions", [])}
    dm = man.get("demo")
    return (dm if isinstance(dm, dict) and dm.get("region") in drift else None), drift


def demo_region(fn: dict) -> dict:
    """Region the demo tour opens (the deck, the speech and DEMO.md talk about the same one).

    best / drift: what the speaker shows live; shot / shot_drift: the region on the screenshots (result.json)."""
    shot = shot_drift = None
    res = SHOTS / "result.json"
    if res.exists():
        try:
            r = json.loads(res.read_text(encoding="utf-8"))
            shot, shot_drift = r.get("best_region"), r.get("drift_region")
        except (OSError, ValueError):
            pass
    if not shot:
        shot, shot_drift = best_region_manifest()
    dm, drift_of = manifest_demo()
    best, best_date, drift_id = shot, None, shot_drift
    if dm:
        best, best_date = dm["region"], dm.get("date")
        has = drift_of.get(best, {}).get(best_date) if best_date else any(drift_of.get(best, {}).values())
        drift_id = best if has else (shot_drift or best_region_manifest()[1])
    regs = {r.get("id"): r for r in (g(fn, "service.regions", []) or [])}

    def info(rid, date=None):
        r = regs.get(rid) or {}
        full = r.get("name") or (rid or DASH)
        date = date or r.get("latest_date")
        dd = f"{date[8:10]}.{date[5:7]}.{date[:4]}" if isinstance(date, str) and len(date) == 10 else DASH
        return {"id": rid, "name": full, "short": full.split(" (")[0], "date": dd,
                "n_det": r.get("n_detections"), "index": r.get("index_permille")}

    return {"best": info(best, best_date), "drift": info(drift_id or best), "shot": info(shot),
            "shot_drift": info(shot_drift or shot), "explicit": bool(dm), "reason": (dm or {}).get("reason", "")}


def numbers(fn: dict) -> dict:
    """All placeholders used in the deck, the speech and the Q&A."""
    l3 = g(fn, "l3_lgbm", {}) or {}
    val = l3.get("val") or {}
    test = l3.get("test") or {}
    rows = g(fn, "models.rows", {}) or {}
    ci = g(fn, "metric_audit.scene_bootstrap_ci95", [None, None]) or [None, None]
    mg_ci = g(fn, "metric_audit.mados_gain_ci95", [None, None]) or [None, None]
    lro = g(fn, "metric_audit.lro", {}) or {}
    per = {r.get("region"): r for r in (lro.get("per_region") or [])}
    dscen = {f"{s.get('region')}": s for s in (g(fn, "drift.scenes", []) or [])}
    hw = str(g(fn, "speed.raw.hardware", "") or "")
    cpu_name = hw.split(",")[0].replace("13th Gen Intel(R) Core(TM) ", "") if hw else DASH
    gpu_name = DASH
    if "NVIDIA GeForce " in hw:
        gpu_name = hw.split("NVIDIA GeForce ")[1].split(" (")[0]
    test_f1 = test.get("f1_md") if isinstance(test, dict) else None
    N = {
        "model": l3.get("model") or "LightGBM",
        "f1": num(val.get("f1_md")), "iou": num(val.get("iou_md")),
        "prec": num(val.get("precision_md")), "rec": num(val.get("recall_md")),
        "rec_pct": pct(val.get("recall_md")), "prec_pct": pct(val.get("precision_md")),
        "rec10": DASH if val.get("recall_md") is None else str(int(round(val["recall_md"] * 10))),
        "thr": num(l3.get("threshold"), 2),
        "n_feat": num(l3.get("n_features"), 0), "n_base": num(l3.get("n_base_features"), 0),
        "n_win": num(l3.get("n_window_features"), 0),
        "train_s": num(l3.get("train_seconds"), 1), "size_mb": num(l3.get("model_size_mb"), 1),
        "seed_std": num(g(fn, "l3_lgbm.noise.f1_std"), 4), "n_seeds": num(g(fn, "l3_lgbm.noise.n_seeds"), 0),
        "rec_low": pct(g(fn, "l3_lgbm.val_recall_by_conf.low")),
        "rec_high": pct(g(fn, "l3_lgbm.val_recall_by_conf.high")),
        "test_computed": bool(g(fn, "test.computed", False)) and test_f1 is not None,
        "test_f1": num(test_f1),
        # models (3-seed means)
        "comb_f1": num(g(rows, "combined.val_f1_mean")), "marida_f1": num(g(rows, "marida_only.val_f1_mean")),
        "unet_f1": num(g(rows, "unet.val_f1_mean")), "stack_f1": num(g(rows, "stack.val_f1_mean")),
        "unet_std": num(g(rows, "unet.val_f1_std"), 4), "comb_std": num(g(rows, "combined.val_f1_std"), 4),
        "unet_thr_acc": num(g(fn, "models.unet_acceptance_threshold")),
        # audit
        "ci_lo": num(ci[0]), "ci_hi": num(ci[1]),
        "thr_opt": num(g(fn, "metric_audit.threshold_optimism"), 4),
        "mados_gain": num(g(fn, "metric_audit.mados_gain")), "mg_lo": num(mg_ci[0]), "mg_hi": num(mg_ci[1]),
        "lro": num(lro.get("mean_f1")), "lro_n": num(lro.get("n_regions"), 0),
        "lro_marida": num(lro.get("mean_f1_marida_only")), "lro_drop": num(lro.get("drop_vs_in_dist")),
        "lro_hon": num(g(per, "honduras_gulf.f1")), "lro_bay": num(g(per, "bay_islands.f1")),
        "lro_haiti": num(g(per, "haiti_pap.f1")),
        # data
        "md_px": num(g(fn, "data.marida.md_px"), 0), "mar_scenes": num(g(fn, "data.marida.n_scenes"), 0),
        "mar_patches": num(g(fn, "data.marida.n_patches"), 0), "mar_regions": num(g(fn, "data.marida.n_regions"), 0),
        "mar_lab_pct": num(g(fn, "data.marida.labelled_pct"), 2),
        "md_med_px": num(g(fn, "data.marida.md_component_median_px"), 0),
        "md_single_pct": num(g(fn, "data.marida.md_single_px_pct"), 0),
        "mados_scenes": num(g(fn, "l3_lgbm.mados.n_scenes"), 0), "mados_patches": num(g(fn, "l3_lgbm.mados.n_patches"), 0),
        "mados_excl": num(g(fn, "l3_lgbm.mados.n_excluded_scenes"), 0),
        "live_n": num(g(fn, "live.n_scenes"), 0), "live_regions": num(g(fn, "live.n_regions"), 0),
        "live_fresh": num(g(fn, "live.n_fresh_2025_2026"), 0),
        # service
        "n_regions": num(g(fn, "service.n_regions"), 0), "n_dates": num(g(fn, "service.n_dates"), 0),
        "n_det": num(g(fn, "service.n_detections_total"), 0), "mdd_thr": num(g(fn, "service.mdd_threshold"), 1),
        "lgbm_map_thr": num(g(fn, "service.lgbm_threshold_map"), 2),
        "live_f1": num(g(fn, "lgbm_live.val.f1_md")),
        "demo_regions": num(g(fn, "service_demo.n_regions"), 0),
        # agreement
        "conf_obj": num(g(fn, "agreement.confirmed_objects"), 0),
        "conf_scenes": num(g(fn, "agreement.scenes_with_confirmed"), 0),
        "agr_scenes": num(g(fn, "agreement.n_scenes"), 0), "agr_r": num(g(fn, "agreement.radius_m"), 0),
        "spearman": num(g(fn, "agreement.spearman_median"), 2),
        # drift
        "drift_n": num(g(fn, "drift.n_scenes"), 0), "drift_regions": num(g(fn, "drift.n_regions"), 0),
        "dc_n": g(fn, "drift_check.n_pairs", DASH), "dc_k": g(fn, "drift_check.k_hit", DASH),
        "dc_kb": g(fn, "drift_check.k_hit_baseline", DASH),
        "drift_h": num(g(fn, "drift.hours"), 0), "drift_cur": g(fn, "drift.currents", DASH),
        "drift_wind": g(fn, "drift.wind", DASH),
        "accra_stranded": num(g(dscen, "accra.stranded_pct"), 0),
        # speed
        "chips": num(g(fn, "speed.raw.chips"), 0), "cuda_s": num(g(fn, "speed.raw.cold_s.cuda"), 1),
        "cpu_s": num(g(fn, "speed.raw.cold_s.cpu"), 1), "base_cuda_s": num(g(fn, "speed.raw.baseline_cold_s.cuda"), 1),
        "warm_gpu": num(g(fn, "speed.raw.warm.cuda.e2e_chips_per_s"), 0),
        "warm_cpu": num(g(fn, "speed.raw.warm.cpu.e2e_chips_per_s"), 1),
        "cpu_name": cpu_name, "gpu_name": gpu_name,
        "ui_load": num(g(fn, "ui_perf.load_s"), 1),
        # speed by CPU threads (affinity), final model
        "gpu_s": num(g(fn, "speed_threads.final.gpu"), 1), "cpu20_s": num(g(fn, "speed_threads.final.cpu20"), 1),
        "cpu8_s": num(g(fn, "speed_threads.final.cpu8"), 1), "cpu4_s": num(g(fn, "speed_threads.final.cpu4"), 1),
        "mid_speedup8": num(g(fn, "rejected.mid.speedup_cpu8"), 2),
        "light_lro": num(g(fn, "rejected.light.lro_mean")), "light_f1": num(g(fn, "rejected.light.val_f1_mean")),
        "z_lro": num(g(fn, "rejected.zfeat.zs_lro_delta")), "z_val": num(g(fn, "rejected.zfeat.zs_val_delta")),
        # confirmed detections over all dates of the service
        "conf_total": num(g(fn, "service.n_confirmed_total"), 0),
        "conf_dates": num(g(fn, "service.n_dates_with_confirmed"), 0),
        # robustness, rehearsal
        "rob_pass": num(g(fn, "robustness.pass"), 0), "rob_n": num(g(fn, "robustness.n_checks"), 0),
        "rob_fail": num(g(fn, "robustness.fail"), 0),
        "reh_first_s": num(g(fn, "rehearsal.first_submit_s"), 0),
        "reh_asis": num(g(fn, "rehearsal.private_asis")), "reh_trained": num(g(fn, "rehearsal.private_trained")),
        "reh_lab": num(g(fn, "rehearsal.labelled_no_seasnot_f1")),
        "reh_human": "–".join(str(x) for x in (g(fn, "rehearsal.human_min", []) or [])) or DASH,
    }
    dr = demo_region(fn)
    N["demo_id"], N["demo_name"], N["demo_short"], N["demo_date"] = (dr["best"]["id"], dr["best"]["name"],
                                                                    dr["best"]["short"], dr["best"]["date"])
    N["demo_ndet"] = num(dr["best"]["n_det"], 0)
    N["demo_index"] = num(dr["best"]["index"], 3)
    N["drift_short"], N["drift_date"] = dr["drift"]["short"], dr["drift"]["date"]
    N["shot_short"], N["shot_date"] = dr["shot"]["short"], dr["shot"]["date"]
    N["shot_drift_short"], N["shot_drift_date"] = dr["shot_drift"]["short"], dr["shot_drift"]["date"]
    N["demo_reason"] = dr["reason"]
    # L43: the drift step is shown in another region when the demo date has no drift
    N["drift_step"] = ("" if dr["drift"]["id"] == dr["best"]["id"]
                       else f"[Левая панель → {N['drift_short']}, снимок {N['drift_date']}] ")
    N["demo_regions_w"] = f"{N['demo_regions']} {plural(g(fn, 'service_demo.n_regions'), 'района', 'районов', 'районов')}"
    if N["test_computed"]:
        N["test_line"] = f"F1 на test MARIDA (посчитан один раз): {N['test_f1']}."
        N["test_speech"] = f"Test MARIDA мы открыли один раз, на финальной модели: F1 {N['test_f1']}."
        N["test_qa"] = (f"Test MARIDA посчитан один раз, на финальной модели, после всех решений: F1 {N['test_f1']}. "
                        "После этого модель не менялась.")
    else:
        N["test_line"] = "Test MARIDA будет посчитан один раз на финальной модели."
        N["test_speech"] = "Test MARIDA мы не открывали: он будет посчитан один раз, на финальной модели."
        N["test_qa"] = ("Test MARIDA будет посчитан один раз на финальной модели; до этого метрики на test не считались, "
                        "и все решения принимались по val.")
    return N


# ------------------------------------------------------------------ speech (one source for SPEECH.md and notes)
# Each entry: slide number, time range, stage directions / text. Placeholders {x} are filled from numbers().
SPEECH = [
    (1, "0:00–0:20",
     "Плавающий мусор в море со спутника — это пятна в один-два пикселя: медианное пятно в разметке MARIDA — "
     "{md_med_px} пикселя по 10 метров. Спутник не взвесит пластик. Но он может подсказать, куда в первую очередь "
     "отправить катер. Кейс просит найти пятна, дать показатель по акватории и сделать это быстро. Мы сделали все три."),
    (2, "0:20–0:40",
     "Данные. Размеченного мусора мало: в MARIDA всего {md_px} пикселей. Поэтому добавили MADOS — ещё {mados_scenes} сцен, "
     "и выкинули из него {mados_excl} сцен, которые пересекаются с проверочной выборкой. Карту строим на {live_n} живых "
     "снимках Sentinel-2 в {live_regions} районах, {live_fresh} из них — 2025–2026 годов."),
    (3, "0:40–1:05",
     "Метод. Пятно — два пикселя, контекста почти нет, решает спектр. Поэтому пиксельный LightGBM: каналы, индексы "
     "и статистики окна вокруг пикселя, всего {n_feat} признаков. UNet мы обучили и сравнили по одному правилу: "
     "{unet_f1} против {comb_f1} у бустинга, UNet отклонён. MADOS принят: плюс {mados_gain} F1, и интервал по сценам "
     "выше нуля. [пауза]"),
    (4, "1:05–1:35",
     "Качество — честно. На проверочной выборке MARIDA, это новые снимки тех же районов, модель находит {rec_pct} "
     "размеченного мусора, и {prec_pct} её находок верны. F1 {f1}. Интервал по сценам — от {ci_lo} до {ci_hi}: сцен мало, "
     "и мы это показываем. Самая строгая проверка — целый регион убран из обучения: F1 {lro}. Столько и ждём на новом "
     "месте. {test_speech}"),
    (5, "1:35–1:50",
     "Показатель на карте — доля видимой воды с признаками мусора, в промилле, по гексагонам H3 около 0,7 км². "
     "Ячейка, видимая меньше чем наполовину, — серая, «нет данных», а не ноль. Это индекс по снимку, не масса пластика. "
     "Покажу вживую. [Alt+Tab в браузер, пауза 2 с]"),
    (6, "1:50–2:20",
     "[Обзор уже на экране] {n_regions_txt}, {n_dates_txt}, справа рейтинг «где искать в первую очередь». "
     "[клик: {demo_short} в надёжном рейтинге] Снимок Sentinel-2 от {demo_date}, кольца — найденные пятна. На живых снимках L2A основной слой — "
     "открытая модель marinedebrisdetector, наша модель — второй слой. [клик по крупному пятну] Карточка находки: "
     "вырезка снимка, площадь помеченной области, уверенность модели."),
    (7, "2:20–2:45",
     "[Слои → Приоритет обследования → клик по зоне №1] Зоны — ответ на вопрос «куда плыть». Каждая объясняет, почему "
     "она первая: пиксели с признаками, умноженные на уверенность, повторяемость по датам и согласие моделей; снимок "
     "с дымкой или бликом получает половину балла, а шов детекторов, кильватер и суда в балл не входят. Двойное кольцо — "
     "уверенная находка: пятно видят две разные модели в радиусе {agr_r} метров, таких {conf_total} по всем датам. Это согласие "
     "моделей, не проверка на месте."),
    (8, "2:45–3:05",
     "{drift_step}[Слои → Дрейф 0→72 ч → ▶] Дрейф на {drift_h} часа: OpenDrift, течения HYCOM, ветер GFS; сиреневое облако — "
     "разброс ветрового сноса. Это демонстрация: на {dc_n} парах снимков прогноз не лучше базовой линии «на месте». [Alt+Tab в презентацию, набрать 9 и Enter]"),
    (9, "3:05–3:35",
     "Скорость. {chips} чипов с холодного старта — {gpu_s} секунды на GPU; на процессоре {cpu20_s} секунды на 20 потоках "
     "и {cpu8_s} на 8 — это честное ограничение: ускорения без потери качества на новом районе мы не нашли. "
     "Ваш датасет подключаем тремя командами. На репетиции с незнакомым архивом первый сабмит ушёл через "
     "{reh_first_s} секунд после распаковки, а обучение на их train подняло скрытый F1 с {reh_asis} до {reh_trained}."),
    (10, "3:35–4:00",
     "Итог. Модель находит {rec_pct} мусора на новых снимках и держит F1 {lro} на невиданном регионе. Карта говорит, "
     "куда плыть, и объясняет почему. Дальше — разметка на L2A, больше дат на район и результаты обследований "
     "обратно в обучение. Спасибо, готовы к вопросам."),
]


def speech_of(n: int, N: dict) -> tuple[str, str]:
    for k, t, s in SPEECH:
        if k == n:
            return t, s.format_map({**N, 'n_regions_txt': f"{N['n_regions']} {plural(N['n_regions'], 'район', 'района', 'районов')}", 'n_dates_txt': f"{N['n_dates']} {plural(N['n_dates'], 'дата', 'даты', 'дат')}"})
    return "", ""


# ------------------------------------------------------------------ drawing helpers
def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def text(slide, x, y, w, h, s, size=20, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         font=FONT, line_spacing=1.1):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    lines = s.split("\n") if isinstance(s, str) else s
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color
    return tb


def rich(slide, x, y, w, h, parts, size=20, align=PP_ALIGN.LEFT, line_spacing=1.15):
    """parts: list of lines; each line a list of (text, color, bold)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    for i, line in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        for t, c, b in line:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.name = FONT
            r.font.bold = b
            r.font.color.rgb = c
    return tb


def rect(slide, x, y, w, h, fill=PANEL, line=None, dash=False, radius=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius:
        shp.adjustments[0] = 0.08
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.5)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    return shp


def header(slide, n, total, title, kicker):
    bg(slide)
    rect(slide, M, Inches(0.5), Inches(0.08), Inches(0.36), fill=CORAL)
    text(slide, M + Inches(0.25), Inches(0.5), Inches(9), Inches(0.35), kicker.upper(), size=14, color=MUTED, bold=True)
    text(slide, M, Inches(1.0), W - 2 * M, Inches(1.2), title, size=30, bold=True, line_spacing=1.0)
    text(slide, W - M - Inches(1.5), H - Inches(0.45), Inches(1.5), Inches(0.3), f"{n} / {total}", size=12,
         color=MUTED, align=PP_ALIGN.RIGHT)
    text(slide, M, H - Inches(0.45), Inches(8), Inches(0.3), "Морской мусор · Sentinel-2", size=12, color=MUTED)


def big_number(slide, x, y, w, value, label, size=88, label_size=20):
    text(slide, x, y, w, Inches(1.4), value, size=size, bold=True, color=CORAL, line_spacing=0.9)
    text(slide, x, y + Inches(1.5), w, Inches(1.6), label, size=label_size, color=MUTED, line_spacing=1.15)


def card(slide, x, y, w, h, title, sub, accent=False, dashed=False, tsize=26, ssize=18):
    rect(slide, x, y, w, h, fill=PANEL, line=CORAL if accent else LINE, radius=True, dash=dashed)
    text(slide, x + Inches(0.22), y + Inches(0.18), w - Inches(0.44), Inches(0.6), title, size=tsize, bold=True,
         color=CORAL if accent else TEXT)
    text(slide, x + Inches(0.22), y + Inches(0.18) + Pt(tsize * 1.45), w - Inches(0.44), h - Inches(0.6), sub,
         size=ssize, color=MUTED, line_spacing=1.1)


def find_shot(names):
    for d in SCREEN_DIRS:
        for nm in names:
            p = d / nm
            if p.exists():
                return p
    for nm in names:  # README images (same frames, JPG) as the last resort
        p = ROOT / "docs" / "img" / IMG_FALLBACK.get(nm, nm)
        if p.exists():
            return p
    return None


def picture(slide, x, y, w, h, names, caption, crop=None, left=False):
    """Screenshot (optionally cropped, px box l,t,r,b of a 1920x1080 frame) fitted into (x,y,w,h-caption)."""
    cap_h = Inches(0.5)
    img = find_shot(names)
    ih = h - cap_h
    if img is not None:
        try:
            from PIL import Image
            with Image.open(img) as im:
                im = im.convert("RGB")
                if crop:
                    sx, sy = im.size[0] / 1920, im.size[1] / 1080
                    im = im.crop((int(crop[0] * sx), int(crop[1] * sy), int(crop[2] * sx), int(crop[3] * sy)))
                iw, ihp = im.size
                buf = io.BytesIO()
                im.save(buf, format="PNG")
            buf.seek(0)
            ar, br = iw / ihp, w / ih
            pw, ph = (w, int(w / ar)) if ar > br else (int(ih * ar), ih)
            px, py = (x if left else x + (w - pw) // 2), y + (ih - ph) // 2
            slide.shapes.add_picture(buf, px, py, pw, ph)
            rect(slide, px, py, pw, ph, fill=None, line=LINE)
            text(slide, px, py + ph + Inches(0.1), max(pw, x + w - px), Inches(0.4), caption, size=16, color=MUTED)
            return True
        except Exception as e:  # noqa: BLE001
            print(f"[deck] cannot place {img}: {e}", file=sys.stderr)
    rect(slide, x, y, w, ih, fill=PANEL, line=LINE, dash=True)
    text(slide, x, y + ih // 2 - Inches(0.3), w, Inches(0.4), caption, size=18, color=MUTED, align=PP_ALIGN.CENTER)
    text(slide, x, y + ih // 2 + Inches(0.2), w, Inches(0.4), " / ".join(names), size=12, color=LINE,
         align=PP_ALIGN.CENTER)
    return False


# ------------------------------------------------------------------ slides
def build(fn: dict, only: int | None = None) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    N = numbers(fn)
    slides = []
    CW = W - 2 * M  # content width
    TOP = Inches(2.45)

    def s1(s, n, t):
        header(s, n, t, "Спутник не взвесит пластик, но подскажет, куда отправить катер", "Задача")
        big_number(s, M, TOP + Inches(0.2), Inches(5.2), f"{N['md_med_px']} пикс.",
                   "медианный размер пятна мусора\nв разметке MARIDA (пиксель 10 м)")
        text(s, Inches(6.4), TOP, Inches(6.3), Inches(0.5), "Кейс просит", size=20, color=MUTED, bold=True)
        items = ["найти пятна мусора на снимке Sentinel-2",
                 "показатель по акватории на карте",
                 "быстрый инференс на данных организаторов"]
        for i, it in enumerate(items):
            yy = TOP + Inches(0.6) + i * Inches(1.05)
            rect(s, Inches(6.4), yy, Inches(6.3), Inches(0.85), fill=PANEL, line=LINE, radius=True)
            text(s, Inches(6.65), yy, Inches(0.5), Inches(0.85), str(i + 1), size=24, bold=True, color=CORAL,
                 anchor=MSO_ANCHOR.MIDDLE)
            text(s, Inches(7.2), yy, Inches(5.4), Inches(0.85), it, size=20, anchor=MSO_ANCHOR.MIDDLE)
    slides.append(s1)

    def s2(s, n, t):
        header(s, n, t, f"Размеченного мусора мало — {N['md_px']} пикселей, поэтому добавили MADOS", "Данные")
        big_number(s, M, TOP + Inches(0.2), Inches(3.6), N["md_px"], "пикселей «мусора»\nв MARIDA (CC BY 4.0)")
        cw, gap = Inches(2.65), Inches(0.2)
        x0 = Inches(4.3)
        card(s, x0, TOP, cw, Inches(3.2), "MARIDA",
             f"{N['mar_scenes']} сцен\n{N['mar_patches']} патчей\nобучение + val\ntest не трогаем")
        card(s, x0 + cw + gap, TOP, cw, Inches(3.2), "MADOS",
             f"+{N['mados_scenes']} сцен\n{N['mados_excl']} сцен убраны:\nпересечение с val", accent=True)
        card(s, x0 + 2 * (cw + gap), TOP, cw, Inches(3.2), "Sentinel-2",
             f"{N['live_n']} живых снимков\n{N['live_regions']} районов\n{N['live_fresh']} — 2025–2026\nбез разметки")
    slides.append(s2)

    def s3(s, n, t):
        header(s, n, t, "Пятно — 2 пикселя, решает спектр: пиксельный LightGBM точнее UNet", "Метод")
        big_number(s, M, TOP + Inches(0.2), Inches(4.2), N["comb_f1"],
                   f"F1 «мусор», val, 3 seed\nLightGBM, {N['n_feat']} признаков:\nканалы, индексы, окна")
        rows = [("LightGBM, MARIDA + MADOS", N["comb_f1"], "принят", True),
                ("стек UNet + LightGBM", N["stack_f1"], "отклонён", False),
                ("LightGBM, только MARIDA", N["marida_f1"], "база", False),
                ("UNet (ResNet-34)", N["unet_f1"], "отклонён", False)]
        extra = {k: v for k, v in (g(fn, "models.rows", {}) or {}).items()
                 if k not in ("combined", "stack", "marida_only", "unet") and isinstance(v, dict)}
        for k, r in list(extra.items())[:1]:  # e.g. a light model added to final_numbers later
            rows.insert(1, (str(r.get("name") or k)[:28], num(r.get("val_f1_mean")),
                            str(r.get("decision") or "")[:10], False))
            rows = rows[:4]
        x0, y0 = Inches(5.2), TOP
        for i, (name, v, dec, acc) in enumerate(rows):
            yy = y0 + i * Inches(0.78)
            rect(s, x0, yy, Inches(7.5), Inches(0.66), fill=PANEL, line=CORAL if acc else LINE, radius=True)
            text(s, x0 + Inches(0.25), yy, Inches(4.4), Inches(0.66), name, size=20, anchor=MSO_ANCHOR.MIDDLE,
                 bold=acc)
            text(s, x0 + Inches(4.6), yy, Inches(1.2), Inches(0.66), v, size=22, bold=True,
                 color=CORAL if acc else TEXT, anchor=MSO_ANCHOR.MIDDLE)
            text(s, x0 + Inches(5.9), yy, Inches(1.5), Inches(0.66), dec, size=18, color=MUTED,
                 anchor=MSO_ANCHOR.MIDDLE)
        text(s, x0, y0 + Inches(3.25), Inches(7.5), Inches(0.8),
             "Правило: принимаем, если прирост F1 ≥ max(0,01; 2·std по seed).", size=18, color=MUTED)
    slides.append(s3)

    def s4(s, n, t):
        header(s, n, t, f"Находим {N['rec_pct']} размеченного мусора на новых снимках тех же районов",
               "Качество, честно")
        big_number(s, M, TOP + Inches(0.2), Inches(4.0), N["f1"],
                   f"F1 «мусор», val MARIDA\nпорог {N['thr']}; точность {N['prec_pct']}")
        cw, gap = Inches(2.65), Inches(0.2)
        x0 = Inches(4.4)  # L39: 3 cards end at the right margin (were 0.2" past it)
        card(s, x0, TOP, cw, Inches(2.2), f"{N['ci_lo']}–{N['ci_hi']}", "95 % интервал\nпо сценам val",
             tsize=24)
        card(s, x0 + cw + gap, TOP, cw, Inches(2.2), N["lro"],
             f"F1 на районе,\nубранном из обучения\n({N['lro_n']} района)", accent=True, tsize=24)
        tt = N["test_f1"] if N["test_computed"] else "test"
        card(s, x0 + 2 * (cw + gap), TOP, cw, Inches(2.2), tt,
             "один раз,\nна финальной модели" if not N["test_computed"] else "MARIDA test,\nпосчитан один раз",
             dashed=not N["test_computed"], tsize=24)
        text(s, x0, TOP + Inches(2.55), Inches(8.3), Inches(1.0),
             f"Решения — только по val. На новом месте ждём около {N['lro']}, а не {N['f1']}.",
             size=20, line_spacing=1.2)
    slides.append(s4)

    def s5(s, n, t):
        header(s, n, t, "Показатель — доля видимой воды с признаками мусора, не масса пластика", "Показатель на карте")
        rect(s, M, TOP, Inches(6.4), Inches(1.5), fill=PANEL, line=CORAL, radius=True)
        text(s, M, TOP, Inches(6.4), Inches(1.5), "‰ = вода с признаками\n     / видимая вода", size=30, bold=True,
             color=CORAL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        rich(s, M, TOP + Inches(1.85), Inches(6.4), Inches(2.4), [
            [("H3 ≈ 0,7 км²", TEXT, True), (" — точка обследования", MUTED, False)],
            [("Видно < 50 %", TEXT, True), (" → «нет данных», не ноль", MUTED, False)],
            [("Индекс по снимку,", CORAL, True), (" не килограммы", MUTED, False)],
        ], size=22, line_spacing=1.6)
        picture(s, Inches(7.3), TOP - Inches(0.1), Inches(5.43), Inches(4.35),
                ["05_h3_2d.png"], f"{N['shot_short']}: индекс по ячейкам H3", crop=(336, 120, 1360, 1040))
    slides.append(s5)

    def s6(s, n, t):
        header(s, n, t, f"{N['n_regions']} {plural(N['n_regions'], 'район', 'района', 'районов')}, {N['n_dates']} {plural(N['n_dates'], 'дата', 'даты', 'дат')}: карта показывает, где искать первым",
               "Живая карта")
        picture(s, M, TOP - Inches(0.35), Inches(7.75), Inches(4.9), ["01_overview_1920.png"],
                f"Обзор: {N['n_regions']} {plural(N['n_regions'], 'район', 'района', 'районов')}, рейтинг «где искать первым»")
        picture(s, M + Inches(7.95), TOP - Inches(0.35), CW - Inches(7.95), Inches(4.9), ["03_detection_card.png"],
                f"Находка: {N['shot_short']}, {N['shot_date']}", crop=(336, 70, 1180, 970))
    slides.append(s6)

    def s7(s, n, t):
        header(s, n, t, "Зона №1 объясняет, почему она первая; двойное кольцо — видят обе модели",
               "Зоны обследования")
        picture(s, M, TOP - Inches(0.15), Inches(5.6), Inches(4.75), ["21_zone_card.png"],
                f"{N['shot_short']}, зона №1: «почему она первая»", crop=(336, 70, 1136, 900), left=True)
        big_number(s, Inches(5.5), TOP + Inches(0.1), Inches(7.2), N["conf_total"],
                   f"пятен видят обе модели в радиусе {N['agr_r']} м\n(на {N['conf_dates']} из {N['n_dates']} {plural(N['n_dates'], 'даты', 'дат', 'дат')})\n"
                   "Это согласие моделей, не проверка на месте", size=88)
    slides.append(s7)

    def s8(s, n, t):
        header(s, n, t, f"Дрейф на {N['drift_h']} ч показывает, куда унесёт пятно — демонстрация без валидации",
               "Дрейф")
        picture(s, M, TOP - Inches(0.15), Inches(5.6), Inches(4.75), ["09_drift.png"],
                f"{N['shot_drift_short']}: дрейф от пятен {N['shot_drift_date']}, T+36 ч", crop=(620, 115, 1540, 1040),
                left=True)
        big_number(s, Inches(5.5), TOP + Inches(0.1), Inches(7.2), N["drift_n"],
                   "снимков с прогнозом дрейфа\nOpenDrift + течения HYCOM + ветер GFS\n"
                   "облако — разброс ветрового сноса 1–3 %", size=88)
    slides.append(s8)

    def s9(s, n, t):
        header(s, n, t, f"{N['chips']} чипов — {N['gpu_s']} с на GPU; первый сабмит на чужих данных — {N['reh_first_s']} с",
               "Скорость и готовность")
        big_number(s, M, TOP + Inches(0.2), Inches(4.2), f"{N['gpu_s']} с",
                   f"GPU {N['gpu_name']}, холодный старт\nCPU: {N['cpu20_s']} с / 20 потоков,\n"
                   f"{N['cpu8_s']} с / 8, {N['cpu4_s']} с / 4")
        x0 = Inches(5.2)
        text(s, x0, TOP, Inches(7.5), Inches(0.45), "Данные организаторов → модель", size=20, bold=True, color=MUTED)
        cmds = [("1", "ingest.py архив.zip", "отчёт «Сомнения»"),
                ("2", "ingest.py архив.zip --convert", "наш формат"),
                ("3", "train_lgbm_ingest.py", f"обучение ≈ {N['train_s']} с")]
        for i, (k, c, d) in enumerate(cmds):
            yy = TOP + Inches(0.6) + i * Inches(0.95)
            rect(s, x0, yy, Inches(7.5), Inches(0.8), fill=PANEL, line=LINE, radius=True)
            text(s, x0 + Inches(0.2), yy, Inches(0.4), Inches(0.8), k, size=22, bold=True, color=CORAL,
                 anchor=MSO_ANCHOR.MIDDLE)
            text(s, x0 + Inches(0.65), yy, Inches(4.6), Inches(0.8), c, size=18, font=MONO,
                 anchor=MSO_ANCHOR.MIDDLE)
            text(s, x0 + Inches(5.1), yy, Inches(2.3), Inches(0.8), d, size=18, color=MUTED,
                 anchor=MSO_ANCHOR.MIDDLE)
        text(s, x0, TOP + Inches(3.5), Inches(7.5), Inches(0.8),
             f"Репетиция: скрытый F1 {N['reh_asis']} «как есть» → {N['reh_trained']} после обучения на их train",
             size=18, color=MUTED)
    slides.append(s9)

    def s10(s, n, t):
        header(s, n, t, "Готово к работе: карта «куда плыть» и модель, которая переобучается за секунды",
               "Выводы и что дальше")
        half = (CW - Inches(0.4)) // 2
        text(s, M, TOP, half, Inches(0.45), "Выводы", size=20, bold=True, color=MUTED)
        text(s, M + half + Inches(0.4), TOP, half, Inches(0.45), "Дальше", size=20, bold=True, color=MUTED)
        left = [(N["lro"], "F1 на новом районе"),
                (f"{N['n_regions']}", "районов на живой карте"),
                (f"{N['gpu_s']} с", f"{N['chips']} чипов на GPU; CPU 8 потоков — {N['cpu8_s']} с")]
        right = ["Разметка на снимках L2A",
                 "Больше дат на район",
                 "Итоги обследований → в обучение"]
        for i, (v, d) in enumerate(left):
            yy = TOP + Inches(0.6) + i * Inches(1.1)
            rect(s, M, yy, half, Inches(0.92), fill=PANEL, line=CORAL if i == 0 else LINE, radius=True)
            text(s, M + Inches(0.25), yy, Inches(1.9), Inches(0.92), v, size=28, bold=True,
                 color=CORAL, anchor=MSO_ANCHOR.MIDDLE)
            text(s, M + Inches(2.2), yy, half - Inches(2.4), Inches(0.92), d, size=20, anchor=MSO_ANCHOR.MIDDLE)
        for i, d in enumerate(right):
            yy = TOP + Inches(0.6) + i * Inches(1.1)
            xx = M + half + Inches(0.4)
            rect(s, xx, yy, half, Inches(0.92), fill=PANEL, line=LINE, radius=True)
            text(s, xx + Inches(0.25), yy, Inches(0.5), Inches(0.92), str(i + 1), size=24, bold=True,
                 color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
            text(s, xx + Inches(0.8), yy, half - Inches(1.0), Inches(0.92), d, size=20, anchor=MSO_ANCHOR.MIDDLE)
    slides.append(s10)

    total = len(slides)
    for i, fn_slide in enumerate(slides, 1):
        if only is not None and i != only:
            continue
        s = prs.slides.add_slide(blank)
        fn_slide(s, i, total)
        tm, sp = speech_of(i, N)
        if sp:
            s.notes_slide.notes_text_frame.text = f"{tm}\n{sp}"
    return prs


# ------------------------------------------------------------------ SPEECH.md
SLIDE_TITLES = {1: "Задача", 2: "Данные", 3: "Метод", 4: "Качество, честно", 5: "Показатель на карте",
                6: "Живая карта (демо)", 7: "Зоны обследования (демо)", 8: "Дрейф (демо)",
                9: "Скорость и готовность", 10: "Выводы и что дальше"}


def speech_md(N: dict) -> str:
    out = [
        "# Речь на защите: 4 минуты дословно",
        "",
        "Файл генерируется `scripts/make_deck.py` из `reports/final_numbers.json`; руками не править, иначе числа "
        "разойдутся с декой. Тот же текст лежит в заметках докладчика `reports/deck.pptx`.",
        "",
        "Как читать: в квадратных скобках — действия и паузы. Слайды 6–8 — живое демо в браузере. "
        "Если демо не поднялось, те же слова говорим по слайдам 6–8 (это скриншоты того же интерфейса), "
        "план Б — `docs/DEMO.md`, §2.",
        "",
        "Темп ~130 слов в минуту. Две репетиции вслух с таймером; если к 1:50 не дошли до демо, "
        "сократите слайд 3 до первой и последней фраз.",
        "",
        "## Текст по слайдам",
        "",
    ]
    for k, t, s in SPEECH:
        out += [f"### {t} — слайд {k}. {SLIDE_TITLES.get(k, '')}", "", s.format_map({**N, 'n_regions_txt': f"{N['n_regions']} {plural(N['n_regions'], 'район', 'района', 'районов')}", 'n_dates_txt': f"{N['n_dates']} {plural(N['n_dates'], 'дата', 'даты', 'дат')}"}), ""]
    out += [
        "## Демо внутри речи (1:50–3:05) и связь с docs/DEMO.md",
        "",
        "| Время | Действие | Соответствует DEMO.md §1 |",
        "|---|---|---|",
        "| 1:50 | Alt+Tab в браузер, обзор уже открыт (http://127.0.0.1:8000, F11) | 0:00–0:30 |",
        f"| 2:00 | Левая панель → {N['demo_short']} ({N['demo_date']}) → клик по крупному пятну | 0:30–1:10 |",
        "| 2:20 | Слои → Приоритет обследования → клик по зоне №1, блок «Почему это место первое» | 1:40–2:10 |",
        f"| 2:45 | {N['drift_step']}Слои → Дрейф 0→72 ч → ▶ | 2:10–2:50 |",
        "| 3:05 | Alt+Tab в презентацию, набрать 9 и Enter | — |",
        "",
        "H3 3D, сравнение районов и вкладку «Проверка» в 4-минутной речи пропускаем; они есть в полном сценарии "
        "DEMO.md §1, если жюри даёт отдельное время на демо или просит показать.",
        "",
        "Если что-то пошло не так — **план Б в `docs/DEMO.md`, §2**: (а) нет интернета → подложка «Без»; "
        "(б) сервис не стартует → `.venv\\Scripts\\python.exe -m service --port 8000`; "
        "(в) всё упало → запасное видео тура или просто листаем слайды 6–8 с тем же текстом. "
        "Потерялись в интерфейсе → «Все районы» в хлебных крошках.",
        "",
        "## Цифры наизусть",
        "",
        "| Число | Что это | Откуда |",
        "|---|---|---|",
        f"| {N['f1']} | F1 класса «мусор» на val MARIDA; порог {N['thr']} | `l3_lgbm.val.f1_md` |",
        f"| {N['rec_pct']} / {N['prec_pct']} | полнота / точность на val | `l3_lgbm.val.recall_md`, `precision_md` |",
        f"| {N['ci_lo']}–{N['ci_hi']} | 95 % интервал F1 по сценам val | `metric_audit.scene_bootstrap_ci95` |",
        f"| {N['lro']} | F1 на районе, убранном из обучения (среднее по {N['lro_n']}) | `metric_audit.lro.mean_f1` |",
        f"| {N['gpu_s']} / {N['cpu20_s']} / {N['cpu8_s']} / {N['cpu4_s']} с | {N['chips']} чипов с холодного старта: GPU / CPU 20 / 8 / 4 потока | `speed_threads.final` |",
        f"| {N['reh_first_s']} с; {N['reh_asis']} → {N['reh_trained']} | репетиция: первый сабмит; скрытый F1 «как есть» → после обучения | `rehearsal` |",
        f"| {N['n_regions']} / {N['n_dates']} | районов / дат на карте | `service.n_regions`, `n_dates` |",
        f"| {N['conf_total']} | пятен, которые видят обе модели в радиусе {N['agr_r']} м (все даты) | "
        "`service.n_confirmed_total` |",
        "",
        f"Про test: {N['test_line']}",
        "",
        "## Самые неприятные вопросы: три готовые фразы",
        "",
        "**«Сколько тонн пластика в этом заливе?»**",
        "",
        "«Мы это сознательно не утверждаем. Пиксель — 10 метров, в нём смесь воды, пены, водорослей и мусора, "
        "а наземных измерений массы для калибровки нет. Наш показатель — доля видимой воды с признаками мусора: "
        "он отвечает на вопрос "
        ""
        "“куда плыть”, а не “сколько килограммов”. С данными обследований его можно откалибровать.»",
        "",
        "**«А на новом месте ваша точность такая же?»**",
        "",
        f"«Нет, и мы это измерили. {N['f1']} — это новые снимки тех же районов. Когда весь район убран из обучения, "
        f"F1 {N['lro']}. Разброс по сценам — от {N['ci_lo']} до {N['ci_hi']}. На вашем датасете переобучим модель "
        "за секунды на вашем train.»",
        "",
        "**«Две модели согласны — значит, это точно мусор?»**",
        "",
        "«Нет. Согласие моделей — не проверка на месте: обе могут ошибиться одинаково, например на фронте мутной воды. "
        "Мы используем согласие только чтобы поднять находку в очереди на проверку человеком и на обследование; "
        "в интерфейсе это так и подписано.»",
        "",
    ]
    return "\n".join(out)


# ------------------------------------------------------------------ qa.md
def qa_items(N: dict) -> list:
    return [
    # --- показатель и честность
    ("Почему вы не оцениваете массу или количество пластика?",
     "Спутник с пикселем 10 м видит не пластик, а спектр поверхности: в пикселе смесь воды, пены, водорослей и мусора, "
     "а наземных данных о массе для калибровки нет. Поэтому показываем индекс по снимку — долю видимой воды с "
     "признаками мусора. Он годится для выбора мест обследования; килограммы по нему утверждать было бы нечестно."),
    ("Что такое ваш показатель и как его читать?",
     "Для ячейки H3 res 8: 1000 × пиксели воды с вероятностью не ниже порога / вся видимая вода, в промилле. "
     "Один пиксель в ячейке — около 0,14 ‰. Рядом всегда дата, модель, порог и observed_frac."),
    ("Что такое observed_frac и зачем серые ячейки?",
     "Доля ячейки, которую удалось увидеть: вода без облаков, теней и пропусков. Меньше 0,5 — ячейка серая, "
     "«нет данных». Ноль под облаком означал бы «чисто», а мы там ничего не видели."),
    ("Почему H3 и почему разрешение 8?",
     "Сетка одинакова для любых дат и районов, сравнение идёт ячейка к ячейке. Ячейка ≈ 0,7 км² — участок, который "
     "реально обследовать за один заход; пикселей в ней достаточно, чтобы доля не прыгала от одного пикселя."),
    ("Как выбираются зоны обследования и почему зона №1 первая?",
     "Базовый балл = пиксели воды с признаками × средняя уверенность модели × бонус повторяемости по надёжным датам; "
     "он умножается на (1 + доля пикселей, подтверждённых второй моделью) и на 0,5, если на снимке дымка или блик. "
     "Шов детекторов, кильватер и суда отфильтрованы как артефакты и в балл не входят. Карточка зоны показывает все "
     "множители и сравнение со следующими зонами; если №1 не подтверждена, плашка «куда отправить» называет и "
     "лучшую подтверждённую зону."),
    # --- метод
    ("Почему LightGBM, а не UNet?",
     f"Пятна крошечные: медиана {N['md_med_px']} пикселя, около {N['md_single_pct']} % пятен — одиночные пиксели. "
     f"Контекста почти нет, решает спектр. UNet мы обучили и сравнили по тому же правилу: {N['unet_f1']} против "
     f"{N['comb_f1']} у бустинга (3 seed), стек тоже не прошёл. Бустинг обучается за секунды на CPU."),
    ("Какие признаки у модели?",
     f"{N['n_feat']} на пиксель: {N['n_base']} — каналы и спектральные индексы (FDI, FAI, NDVI, NDWI и др.), "
     f"{N['n_win']} — статистики окон 3–31 пикселя (среднее, разброс, отличие от локальной медианы). Окна дают "
     "контраст пятна с окружающей водой."),
    ("Зачем MADOS и не утечка ли это?",
     f"В MARIDA всего {N['md_px']} пикселей мусора. MADOS дал +{N['mados_gain']} F1, интервал по сценам "
     f"{N['mg_lo']}–{N['mg_hi']} — выше нуля. {N['mados_excl']} сцен MADOS, совпадающих по месту или съёмке с val "
     "MARIDA, из обучения убраны; сцены, совпадающие с test MARIDA, тоже не используются."),
    ("Почему порог именно " + N["thr"] + "?",
     f"Это argmax F1 на val MARIDA по сетке 0,02–0,98. Кривая вокруг оптимума плоская. Оптимизм от выбора порога "
     f"на тех же данных мы измерили: около {N['thr_opt']} F1. На живых снимках порог показа свой и всегда подписан."),
    # --- качество и валидация
    ("Как вы валидировали?",
     "Все решения — признаки, выборка, порог, модель — только по официальному val MARIDA. Шум — повтор с 3 seed; "
     "изменение принимаем, если прирост F1 не меньше max(0,01; 2·std). " + N["test_qa"]),
    ("Что значит интервал по сценам?",
     f"В val всего несколько сцен, и F1 сильно зависит от того, какие попали. Бутстрэп по сценам даёт 95 % "
     f"интервал {N['ci_lo']}–{N['ci_hi']}. Разброс ±{N['seed_std']} по seed — это только шум обучения, не "
     "неопределённость метрики."),
    ("Нет ли утечки «того же места» между train и val?",
     "Общих патчей и сцен нет, но часть val лежит в тех же районах, что train, другой датой. Мы разделили val на "
     "«перекрывается с train» и «не перекрывается»: разница не отличима от нуля. Поэтому и честно называем val "
     "«новые снимки тех же районов»."),
    ("Что будет на совсем новом месте?",
     f"Проверили leave-region-out: регион целиком убран из обучения, порог выбран без него. Средний F1 "
     f"{N['lro']} по {N['lro_n']} регионам (от {N['lro_bay']} до {N['lro_haiti']}); без MADOS было бы "
     f"{N['lro_marida']}. Это наш прогноз для нового места."),
    ("Почему test не посчитан?",
     N["test_qa"] + " Так метрика test остаётся независимой от наших решений."),
    ("Почему не оценили marinedebrisdetector на MARIDA?",
     "Он обучался в том числе на MARIDA, любая цифра была бы завышенной. Используем его как готовый открытый "
     "инструмент для живых снимков L2A (MIT) и не выдаём его метрики за свои."),
    # --- живые снимки и ложные срабатывания
    ("Модель обучена на ACOLITE, а живые снимки — L2A. Это не проблема?",
     f"Проблема, и мы её называем. Уровни отражения разные, вероятность на L2A не откалибрована. Поэтому на живых "
     f"снимках основной слой — marinedebrisdetector, обученный на L2A, а наша модель (вариант с L2A-аугментацией, "
     f"F1 val {N['live_f1']}) — второй слой и подтверждение."),
    ("Две модели согласны — значит, это мусор?",
     f"Нет. Согласие — не проверка на месте: обе модели могут ошибаться одинаково, например на фронте мутной воды. "
     f"Уверенных пятен {N['conf_total']} на {N['conf_dates']} из {N['n_dates']} дат; по пикселям модели "
     f"согласуются слабо (медианный ρ Спирмена {N['spearman']}). Согласие только поднимает находку в очереди на проверку."),
    ("Какие ложные срабатывания главные?",
     "Природная органика (водоросли, Sargassum), суда и их следы, пена и прибой, мутные речные шлейфы, кромки облаков, "
     "блик. В обучении есть все эти классы как трудные негативы; на карте — маски облаков и суши, отбрасывание "
     "одиночных пикселей, карточка с вырезкой снимка для глаз человека."),
    ("А суда?",
     "Судно на снимке 10 м похоже на яркое пятно. В MARIDA и MADOS есть класс «судно», модель учится его отличать, "
     "но на val часть судов ещё путается с мусором. В карточке находки видна вырезка: судно с кильватерным следом "
     "человек отличает за секунду и отмечает «ложное»."),
    ("Что с дымкой и солнечным бликом?",
     "Блик поднимает ИК-каналы — это ровно сигнатура мусора, поэтому это главный риск. Снимки выбираем с минимальным "
     "бликом, снимки с сильной дымкой отбраковываем, а оставшиеся помечаем жёлтой плашкой «дымка» и в календаре "
     "как «ненадёжно». Синтетические проверки облаков и блика — в reports/robustness.md."),
    ("Что будет на облачном снимке?",
     "Облака и тени исключаются маской, такие ячейки серые «нет данных». Облачность видна в KPI и в observed_frac; "
     "для района лучше взять другую дату — календарь показывает, какие даты надёжны."),
    ("Что такое «проверка человеком» и когда дообучение меняет модель?",
     "Во вкладке «Проверка» оператор помечает находки клавишами 1–6, кнопка «Дообучить» переобучает LightGBM с его "
     "метками и сравнивает F1 на val MARIDA до и после. Правило то же: прирост ≥ 0,01. Пока прирост меньше — "
     "решение «не принято», веса не подменяются: метки на L2A немного меняют val на ACOLITE."),
    # --- дрейф
    ("Что с прогнозом дрейфа?",
     f"Это демонстрация, не прогноз: OpenDrift, течения {N['drift_cur']}, ветер {N['drift_wind']}, "
     f"{N['drift_h']} ч, ансамбль ветрового сноса 1–3 %. Проверили на {N['dc_n']} парах повторных снимков (эксперимент, "
     f"критерий записан заранее): прогноз {N['dc_k']}/{N['dc_n']}, базовая линия «на месте» {N['dc_kb']}/{N['dc_n']} — не лучше. "
     "Причины: вырезка 25 км мала, в бухтах частицы уходят на берег, находки второго снимка — не те же объекты. "
     "Для старых дат дрейф не считаем."),
    ("HYCOM у берега — это же 1/12°?",
     f"Да, ячейка около 9 км, у берега течений нет или они грубые; пропуски заполняются ближайшими значениями. "
     f"Частицы, дошедшие до берега, остаются на месте выброса — в Аккре так закончили {N['accra_stranded']} % частиц. "
     "Поэтому у берега дрейф показывает только направление, а не траекторию."),
    # --- скорость и данные организаторов
    ("Насколько это быстро?",
     f"{N['chips']} чипов 256×256 с холодного старта: {N['gpu_s']} с на GPU {N['gpu_name']}, {N['cpu20_s']} с на CPU "
     f"({N['cpu_name']}, 20 потоков); до ускорения на GPU было {N['base_cuda_s']} с. Выход бит в бит тот же "
     f"(тест эквивалентности). Карта грузится за {N['ui_load']} с."),
    ("А на чужом железе без NVIDIA?",
     f"CPU-путь работает без CUDA и выбирается автоматически. Время почти пропорционально числу потоков: "
     f"{N['cpu8_s']} с на 8 потоках и {N['cpu4_s']} с на 4 — это наше честное ограничение. Лёгкая модель в ~3 раза быстрее "
     f"при том же F1 на val ({N['light_f1']}), но на новом районе хуже ({N['light_lro']} против {N['lro']}), поэтому не "
     f"принята; средняя почти не быстрее (×{N['mid_speedup8']} на 8 потоках): время уходит на предсказание леса."),
    ("Что если у вас будут только RGB-каналы?",
     "Модель переобучается на доступных каналах за секунды. Без B1 потерь нет, 10 м + SWIR — почти без потерь, "
     "RGB+NIR — заметно хуже, только RGB — хуже всего: без ИК нет FDI и контраста пятна в B8. Оценки по val и "
     "команды — reports/l23_channels_speed.md."),
    ("Как вы подключите наш датасет?",
     "Три команды: ingest.py архив.zip (безопасная распаковка, HTML-отчёт с разделом «Сомнения», черновик конфига), "
     "то же с --convert (каналы, масштаб, классы → наш формат), train_lgbm_ingest.py (обучение, порог по вашему val). "
     "Проверено тестами на трёх искусственных наборах разных форматов; пересечения с MARIDA/MADOS проверяет provenance_check.py."),
    # --- лицензии и воспроизводимость
    ("Что показала репетиция на незнакомом датасете?",
     f"Архив, упакованный как у организаторов: первый валидный сабмит через {N['reh_first_s']} с машинного времени после "
     f"распаковки (человеку {N['reh_human']} мин). Скрытый F1 «как есть» {N['reh_asis']}, после обучения на их train — "
     f"{N['reh_trained']}. Разобрали, почему «как есть» низкий: инструменты не искажают данные, на размеченных пикселях без "
     f"класса sea snot F1 {N['reh_lab']}; всё съедают метрика по неразмеченным пикселям и другое определение негатива. "
     "Вывод: порог подбираем под их метрику, их «прочее» добавляем в негативы."),
    ("Какие лицензии у данных и моделей?",
     "MARIDA и MADOS — CC BY 4.0; marinedebrisdetector — MIT; снимки Sentinel-2 — открытая лицензия Copernicus "
     "(«Contains modified Copernicus Sentinel data»); подложки CARTO/Esri и OpenStreetMap — с атрибуцией на карте. "
     "Полный список — reports/sources.md."),
    ("Как воспроизвести одной командой?",
     "powershell -ExecutionPolicy Bypass -File run.ps1: создаёт окружение, поднимает сервис и открывает карту. "
     "Инференс: python inference.py --data-dir <папка> --output <папка>. Числа README, отчёта и этой деки "
     "генерируются скриптами из reports/final_numbers.json; seed фиксированы, веса LightGBM лежат в репозитории. "
     f"В чистом клоне — демо-набор из {N['demo_regions_w']}; полный слой карты подключается через -DataRoot."),
]


def qa_md(N: dict) -> str:
    out = ["# Вопросы жюри: короткие честные ответы", "",
           "Файл генерируется `scripts/make_deck.py` из `reports/final_numbers.json`; числа руками не править. "
           f"{N['test_line']}", ""]
    for i, (q, a) in enumerate(qa_items(N), 1):
        out += [f"**{i}. {q}**", a, ""]
    return "\n".join(out)


# ------------------------------------------------------------------ preview / main
def find_soffice():
    for c in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if c and Path(c).exists():
            return c
    return None


def preview(fn: dict, out_dir: Path, n_slides: int) -> int:
    so = find_soffice()
    if not so:
        print("[deck] LibreOffice (soffice) not found: preview skipped")
        return 0
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("slide_*.png"):
        old.unlink()
    made = 0
    with tempfile.TemporaryDirectory() as td:
        prof = Path(td) / "lo_profile"
        for i in range(1, n_slides + 1):
            p = Path(td) / f"slide_{i:02d}.pptx"
            build(fn, only=i).save(p)
            cmd = [so, f"-env:UserInstallation=file:///{prof.as_posix()}", "--headless", "--convert-to", "png",
                   "--outdir", str(out_dir), str(p)]
            try:
                subprocess.run(cmd, check=True, capture_output=True, timeout=180)
                made += (out_dir / f"slide_{i:02d}.png").exists()
            except Exception as e:  # noqa: BLE001
                print(f"[deck] preview slide {i} failed: {e}")
    print(f"[deck] preview: {made}/{n_slides} PNG -> {out_dir}")
    return made


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="deck.pptx + SPEECH.md + qa.md из final_numbers.json")
    ap.add_argument("--numbers", default=str(ROOT / "reports" / "final_numbers.json"))
    ap.add_argument("--out", default=str(ROOT / "reports" / "deck.pptx"))
    ap.add_argument("--preview", action="store_true", help="PNG per slide via LibreOffice -> reports/deck_preview")
    ap.add_argument("--no-docs", action="store_true", help="do not rewrite docs/SPEECH.md and reports/qa.md")
    a = ap.parse_args(argv)
    p = Path(a.numbers)
    fn = json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}
    if not fn:
        print(f"[deck] {p} not found: numbers will be dashes (run scripts/final_numbers.py)")
    prs = build(fn)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    n = len(prs.slides)
    print(f"[deck] {n} slides -> {out}")
    if not a.no_docs:
        N = numbers(fn)
        (ROOT / "docs" / "SPEECH.md").write_text(speech_md(N), encoding="utf-8")
        (ROOT / "reports" / "qa.md").write_text(qa_md(N), encoding="utf-8")
        print(f"[deck] docs/SPEECH.md, reports/qa.md ({len(qa_items(N))} вопросов) written")
    if a.preview:
        preview(fn, ROOT / "reports" / "deck_preview", n)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
