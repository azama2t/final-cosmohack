r"""Build the pitch deck reports/deck.pptx (12 slides, 16:9, dark theme) from reports/final_numbers.json.

Usage (from repo root):
    .venv\Scripts\python.exe scripts\make_deck.py                 # -> reports\deck.pptx
    .venv\Scripts\python.exe scripts\make_deck.py --preview       # + PNG per slide in reports\deck_preview\ (LibreOffice)

Numbers come only from reports/final_numbers.json (missing -> "—"). Screenshots: docs/img/*.jpg (in git), else
reports/screens/iter6/*.png (if absent, a framed placeholder with the expected file name is drawn).
Rules: the title states the conclusion; one big number per slide; <= 25 words of body text; captions on images.
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCREEN_DIRS = [ROOT / "docs" / "img", ROOT / "reports" / "screens" / "iter6"]
SCREENS = SCREEN_DIRS[0]

BG = RGBColor(0x0A, 0x11, 0x1F)       # deep blue-black
PANEL = RGBColor(0x12, 0x1C, 0x2E)
LINE = RGBColor(0x24, 0x33, 0x4A)
TEXT = RGBColor(0xE8, 0xEE, 0xF6)
MUTED = RGBColor(0x8C, 0x9D, 0xB5)
CORAL = RGBColor(0xFF, 0x6B, 0x5A)    # the single accent
FONT = "Segoe UI"  # Inter is not installed on the demo PC; Segoe UI is present on every Windows
DASH = "—"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)


# ------------------------------------------------------------------ helpers
def g(d, path, default=None):
    cur = d
    for p in path.split("."):
        if isinstance(cur, dict) and p in cur:
            cur = cur[p]
        elif isinstance(cur, list):
            try:
                cur = cur[int(p)]
            except (ValueError, IndexError):
                return default
        else:
            return default
    return default if cur is None else cur


def f(v, nd=2, pct=False):
    if v is None:
        return DASH
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    if pct:
        return f"{100 * x:.0f} %"
    if nd == 0:
        return f"{int(round(x)):,}".replace(",", " ")
    return f"{x:.{nd}f}"


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def text(slide, x, y, w, h, s, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         font=FONT, line_spacing=1.1):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    lines = s.split("\n") if isinstance(s, str) else s
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill=PANEL, line=None, dash=False, radius=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius:
        shp.adjustments[0] = 0.08
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    return shp


def header(slide, n, total, title, kicker):
    bg(slide)
    rect(slide, M, Inches(0.55), Inches(0.08), Inches(0.42), fill=CORAL)
    text(slide, M + Inches(0.25), Inches(0.52), Inches(9), Inches(0.3), kicker.upper(), size=12, color=MUTED, bold=True)
    text(slide, M, Inches(1.05), W - 2 * M, Inches(1.3), title, size=32, bold=True, line_spacing=1.0)
    text(slide, W - M - Inches(1.5), H - Inches(0.5), Inches(1.5), Inches(0.3), f"{n} / {total}", size=11,
         color=MUTED, align=PP_ALIGN.RIGHT)
    text(slide, M, H - Inches(0.5), Inches(8), Inches(0.3), "Макропластик · Sentinel-2", size=11, color=MUTED)


def big_number(slide, x, y, w, value, label, size=96):
    text(slide, x, y, w, Inches(1.6), value, size=size, bold=True, color=CORAL, line_spacing=0.9)
    text(slide, x, y + Inches(1.65), w, Inches(0.9), label, size=18, color=MUTED)


def picture(slide, x, y, w, h, name_hint, caption, placeholder_name):
    """Screenshot fitted into (x,y,w,h) with a caption bar; placeholder frame if no file."""
    img = None
    for d in SCREEN_DIRS:
        if img is not None or not d.exists():
            continue
        cands = sorted([*d.glob("*.jpg"), *d.glob("*.png")])
        for k in name_hint:  # hints in priority order
            hit = next((c for c in cands if k in c.stem.lower()), None)
            if hit is not None:
                img = hit
                break
    if img is not None:
        try:
            from PIL import Image
            with Image.open(img) as im:
                iw, ih = im.size
            ar, br = iw / ih, w / h
            if ar > br:
                pw, ph = w, int(w / ar)
            else:
                pw, ph = int(h * ar), h
            px, py = x + (w - pw) // 2, y + (h - ph) // 2
            slide.shapes.add_picture(str(img), px, py, pw, ph)
            rect(slide, px, py + ph - Inches(0.45), pw, Inches(0.45), fill=BG)
            text(slide, px + Inches(0.15), py + ph - Inches(0.38), pw - Inches(0.3), Inches(0.3), caption,
                 size=13, color=TEXT)
            return True
        except Exception as e:  # noqa: BLE001
            print(f"[deck] cannot place {img}: {e}", file=sys.stderr)
    rect(slide, x, y, w, h, fill=PANEL, line=LINE, dash=True)
    text(slide, x, y + h // 2 - Inches(0.4), w, Inches(0.4), caption, size=16, color=MUTED, align=PP_ALIGN.CENTER)
    text(slide, x, y + h // 2 + Inches(0.1), w, Inches(0.4), f"docs/img/{placeholder_name}",
         size=12, color=LINE, align=PP_ALIGN.CENTER)
    return False


def chips(slide, x, y, items, w_each, h=Inches(1.1), gap=Inches(0.2), accent_idx=None, size=16):
    for i, (t, sub) in enumerate(items):
        xx = x + i * (w_each + gap)
        rect(slide, xx, y, w_each, h, fill=PANEL, line=CORAL if i == accent_idx else LINE, radius=True)
        text(slide, xx + Inches(0.2), y + Inches(0.18), w_each - Inches(0.4), Inches(0.4), t, size=size, bold=True)
        if sub:
            text(slide, xx + Inches(0.2), y + Inches(0.6), w_each - Inches(0.4), Inches(0.4), sub, size=12, color=MUTED)


# ------------------------------------------------------------------ slides
def build(fn: dict, only: int | None = None) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    slides = []

    l3 = g(fn, "l3_lgbm", {}) or {}
    val = l3.get("val") or {}
    test = l3.get("test") or {}
    mar = g(fn, "data.marida", {}) or {}
    sv = g(fn, "service", {}) or {}
    models = (g(fn, "models.rows", {}) or {})
    lro = g(fn, "metric_audit.lro", {}) or {}
    agr = g(fn, "agreement", {}) or {}
    drift = g(fn, "drift", {}) or {}

    # 1. Problem
    def s1(s, n, t):
        header(s, n, t, "Мусор в море виден со спутника как сигнал, куда отправить обследование", "Проблема")
        big_number(s, M, Inches(2.9), Inches(5), "10 м", "пиксель Sentinel-2: пятно мусора\nзанимает 1–4 пикселя")
        text(s, Inches(6.6), Inches(3.0), Inches(6.1), Inches(3),
             "Судов и дронов мало, акватории огромные.\nНужна карта: где подозрительно, насколько, куда плыть первым.",
             size=22, line_spacing=1.25)
    slides.append(s1)

    # 2. Data
    def s2(s, n, t):
        md = f(mar.get("md_px"), 0)
        header(s, n, t, f"Мусора в MARIDA всего {md} размеченных пикселей: ищем иголки в стоге", "Данные")
        big_number(s, M, Inches(2.9), Inches(5.5), md, "пикселей Marine Debris\n(High / Moderate / Low уверенность)")
        chips(s, Inches(6.6), Inches(3.0), [
            (f(mar.get("n_patches"), 0), "патчей 256×256"),
            (f(mar.get("n_scenes"), 0), "сцен Sentinel-2"),
            (f"{mar.get('labelled_pct') if mar.get('labelled_pct') is not None else DASH} %", "пикселей размечено"),
        ], Inches(1.9))
        text(s, Inches(6.6), Inches(4.5), Inches(6.1), Inches(1.2),
             "15 классов, включая похожие на мусор: пена, Sargassum, суда, следы.\nЛицензия CC BY 4.0.",
             size=16, color=MUTED, line_spacing=1.25)
    slides.append(s2)

    # 3. Method
    def s3(s, n, t):
        header(s, n, t, "Пятна по 2 пикселя: решает спектр, поэтому пиксельный бустинг", "Метод")
        big_number(s, M, Inches(2.9), Inches(3.5), f(l3.get("n_features"), 0),
                   f"признаков на пиксель:\n{f(l3.get('n_base_features'), 0)} спектр + {f(l3.get('n_window_features'), 0)} окно")
        chips(s, Inches(4.4), Inches(3.1), [
            ("Данные", "MARIDA + MADOS"), ("Признаки", "каналы, индексы, окна"), ("LightGBM", "мусор / не мусор"),
            ("Порог", f"{f(l3.get('threshold'))} по val"), ("H3", "показатель, зоны"),
        ], Inches(1.55), gap=Inches(0.15), accent_idx=2, size=15)
        text(s, Inches(4.4), Inches(4.6), Inches(8.3), Inches(1.2),
             "На живых L2A-сценах основной слой — marinedebrisdetector (MIT); на MARIDA его не оцениваем.",
             size=16, color=MUTED, line_spacing=1.25)
    slides.append(s3)

    # 4. Index
    def s4(s, n, t):
        header(s, n, t, "Показатель: какая доля видимой воды выглядит как мусор", "Показатель")
        rect(s, M, Inches(2.8), W - 2 * M, Inches(1.6), fill=PANEL, radius=True)
        text(s, M, Inches(3.05), W - 2 * M, Inches(1.2), "‰ = помеченная вода / видимая вода",
             size=38, bold=True, color=CORAL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, M, Inches(4.8), Inches(6), Inches(1.4),
             "Ячейка H3 ≈ 0,7 км².\nobserved_frac < 0,5 → серая «нет данных», не ноль.", size=18, line_spacing=1.3)
        rect(s, Inches(7.2), Inches(4.8), Inches(5.5), Inches(0.9), fill=None, line=CORAL, radius=True)
        text(s, Inches(7.4), Inches(4.95), Inches(5.1), Inches(0.6), "Индекс по снимку, не масса пластика",
             size=20, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slides.append(s4)

    # 5. Quality
    def s5(s, n, t):
        rec = val.get("recall_md")
        prec = val.get("precision_md")
        k = DASH if rec is None else str(int(round(rec * 10)))
        kp = DASH if prec is None else str(int(round(prec * 10)))
        header(s, n, t, f"Модель находит {k} из 10 пикселей мусора, и {kp} из 10 её находок верны", "Качество (MARIDA)")
        big_number(s, M, Inches(2.9), Inches(4.5), f(val.get("f1_md"), 3), "F1 Marine Debris, val\n(выбор модели и порога)")
        chips(s, Inches(5.6), Inches(3.0), [
            (f(val.get("iou_md"), 3), "IoU val"),
            (f(lro.get("mean_f1"), 3), "F1 на новом районе"),
            ("0.80", "RF из статьи MARIDA, test"),
        ], Inches(2.2), accent_idx=1)
        test_txt = (f"F1 test (один раз): {f(test.get('f1_md'))}." if test.get("f1_md") is not None
                    else "Test посчитаем один раз на итоговой модели.")
        text(s, Inches(5.6), Inches(4.5), Inches(7.1), Inches(1.4),
             f"Выбор только по val, шум — 3 seed. {test_txt}", size=16, color=MUTED, line_spacing=1.25)
    slides.append(s5)

    # 5b. Model comparison
    def s5b(s, n, t):
        header(s, n, t, "MADOS принят, UNet и стек отклонены по одному правилу", "Сравнение моделей")
        items = []
        for key, label in (("marida_only", "только MARIDA"), ("combined", "MARIDA + MADOS"),
                           ("unet", "UNet"), ("stack", "стек UNet + LGBM")):
            r = models.get(key) or {}
            m_, sd = r.get("val_f1_mean"), r.get("val_f1_std")
            val_s = DASH if m_ is None else f"{m_:.3f}"
            sub = label + ("" if sd is None else f" · ±{sd:.4f}")
            items.append((val_s, sub))
        chips(s, M, Inches(3.0), items, Inches(2.85), h=Inches(1.3), accent_idx=1, size=28)
        text(s, M, Inches(4.7), W - 2 * M, Inches(1.2),
             "F1 Marine Debris на val, 3 seed. Принимаем, если прирост ≥ max(0.01, 2×std).",
             size=16, color=MUTED, line_spacing=1.25)
    slides.append(s5b)

    # 6. Live map
    def s6(s, n, t):
        header(s, n, t, "Живая карта: снимок, находки и зоны на одном экране", "Демо")
        picture(s, M, Inches(2.75), Inches(8.6), Inches(4.1), ["region", "02_region"], "Снимок + находки модели",
                "region.jpg")
        big_number(s, Inches(9.6), Inches(3.0), Inches(3.2), f(sv.get("n_regions"), 0),
                   f"районов, {f(sv.get('n_dates'), 0)} дат\nсцен Sentinel-2 L2A", size=88)
    slides.append(s6)

    # 7. Zones & compare
    def s7(s, n, t):
        header(s, n, t, "Топ ячеек H3 — готовый список точек для выхода судна", "Зоны и сравнение")
        picture(s, M, Inches(2.75), Inches(8.6), Inches(4.1), ["07_zones", "zones", "compare"],
                "Приоритет обследования и сравнение районов", "07_zones.png")
        big_number(s, Inches(9.6), Inches(3.0), Inches(3.2), "10",
                   "зон на район: ранг,\nкоординаты, причина", size=88)
    slides.append(s7)

    # 7b. Confirmed detections
    def s7b(s, n, t):
        header(s, n, t, "Уверенная находка: две разные модели видят объект в пределах 20 м", "Согласие моделей")
        picture(s, M, Inches(2.75), Inches(8.6), Inches(4.1), ["detection_card", "12_card_confirmed"],
                "Карточка находки: вырезка снимка и согласие моделей", "detection_card.jpg")
        big_number(s, Inches(9.6), Inches(3.0), Inches(3.2), f(agr.get("confirmed_objects"), 0),
                   f"объектов в {f(agr.get('scenes_with_confirmed'), 0)} из {f(agr.get('n_scenes'), 0)}\nсцен; не проверка на месте",
                   size=88)
    slides.append(s7b)

    # 8. Drift
    def s8(s, n, t):
        header(s, n, t, "Дрейф на 72 часа: демонстрация направления, не прогноз", "Дрейф")
        picture(s, M, Inches(2.75), Inches(8.6), Inches(4.1), ["drift"], "Треки частиц 0–72 ч (без валидации)",
                "drift.jpg")
        big_number(s, Inches(9.6), Inches(3.0), Inches(3.2), f(drift.get("n_scenes"), 0),
                   "свежих сцен с дрейфом:\nOpenDrift, HYCOM, GFS", size=88)
    slides.append(s8)

    # 9. Tools
    def s9(s, n, t):
        header(s, n, t, "Новый датасет — первый сабмит за 60 минут", "Готовность к данным организаторов")
        big_number(s, M, Inches(2.9), Inches(4), "60", "минут от архива\nдо первого сабмита")
        chips(s, Inches(4.9), Inches(3.0), [
            ("inspect", "каналы, масштаб, классы"),
            ("forensics", "разметка из спектра?"),
            ("adapter", "YAML → наш формат"),
            ("provenance", "пересечения с MARIDA"),
        ], Inches(1.8))
        text(s, Inches(4.9), Inches(4.5), Inches(7.8), Inches(1.2),
             "Проверено на MARIDA и чужом датасете Sentinel-2. Дальше — переобучение LightGBM за минуты.",
             size=16, color=MUTED, line_spacing=1.25)
    slides.append(s9)

    # 10. Next
    def s10(s, n, t):
        header(s, n, t, "Дальше: калибровка на L2A и повторяемость находок во времени", "Что дальше")
        items = [("1", "Дообучить модель на размеченных L2A-сценах"),
                 ("2", "Больше дат на район: повторяемость = надёжность"),
                 ("3", "Результаты обследований — обратно в разметку")]
        for i, (num, body) in enumerate(items):
            y = Inches(2.8) + i * Inches(1.2)
            text(s, M, y, Inches(1), Inches(1), num, size=48, bold=True, color=CORAL if i == 0 else MUTED)
            text(s, M + Inches(1.1), y + Inches(0.2), Inches(10.5), Inches(0.8), body, size=24)
    slides.append(s10)

    total = len(slides)
    for i, fn_slide in enumerate(slides, 1):
        if only is not None and i != only:
            continue
        s = prs.slides.add_slide(blank)
        fn_slide(s, i, total)
    return prs


def find_soffice():
    for c in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if c and Path(c).exists():
            return c
    return None


def preview(fn: dict, out_dir: Path, n_slides: int) -> int:
    so = find_soffice()
    if not so:
        print("[deck] LibreOffice (soffice) not found: preview skipped")
        return 0
    out_dir.mkdir(parents=True, exist_ok=True)
    made = 0
    with tempfile.TemporaryDirectory() as td:
        prof = Path(td) / "lo_profile"
        for i in range(1, n_slides + 1):
            p = Path(td) / f"slide_{i:02d}.pptx"
            build(fn, only=i).save(p)
            cmd = [so, f"-env:UserInstallation=file:///{prof.as_posix()}", "--headless", "--convert-to", "png",
                   "--outdir", str(out_dir), str(p)]
            try:
                subprocess.run(cmd, check=True, capture_output=True, timeout=120)
                made += (out_dir / f"slide_{i:02d}.png").exists()
            except Exception as e:  # noqa: BLE001
                print(f"[deck] preview slide {i} failed: {e}")
    print(f"[deck] preview: {made}/{n_slides} PNG -> {out_dir}")
    return made


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="reports/deck.pptx из final_numbers.json")
    ap.add_argument("--numbers", default=str(ROOT / "reports" / "final_numbers.json"))
    ap.add_argument("--out", default=str(ROOT / "reports" / "deck.pptx"))
    ap.add_argument("--preview", action="store_true", help="PNG per slide via LibreOffice -> reports/deck_preview")
    a = ap.parse_args(argv)
    p = Path(a.numbers)
    fn = json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}
    if not fn:
        print(f"[deck] {p} not found: numbers will be dashes (run scripts/final_numbers.py)")
    prs = build(fn)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    n = len(prs.slides)
    shots = sorted(x.name for x in SCREENS.glob("*.jpg")) if SCREENS.exists() else []
    print(f"[deck] {n} slides -> {out}; screenshots used from {SCREENS}: {shots or 'none (placeholders)'}")
    if a.preview:
        preview(fn, ROOT / "reports" / "deck_preview", n)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
