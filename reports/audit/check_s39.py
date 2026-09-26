"""Audit L107 (INBOX §39 п.1–3): statuses, classification, quantity line — API, CSV/GeoJSON export, UI card/list/legend, docs, deck.
  S1 (critical) status of a zone not one of the 4 statement statuses; «требует проверки» used as a status anywhere
     (API status_label/detection_label, CSV status, UI status chip / list, deck, docs);
  S2 (important) confirmation: level-B zones «совпадает с разметкой Cózar 2024 (B)», other finds «независимой разметки нет»;
  S3 (critical) «Исключено» lists a background whose check was NOT run or whose flag is raised (e.g. «ветер» excluded but wind
     unknown / >= 5 m/s; «судно» excluded but ship flag true); (important) algae / sargassum not in «Не проверяется»;
  S4 (critical) quantity: first line of «Количество» in the card is not «Количество предметов по этому снимку не определено»;
     items/km2 visible without expanding; items/km2 in the zone list; (important) «нижняя граница» anywhere;
  S5 (important) export: status has 4 values; confirmation, class, excluded_backgrounds, scenario_status present.
Usage: .venv/Scripts/python.exe reports/audit/check_s39.py http://127.0.0.1:8070
Output -> reports/audit/s39.json + ui_s39_*.png."""
import csv
import io
import json
import re
import sys
import time
import urllib.request
from pathlib import Path

from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "reports" / "audit"
B = (sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8070").rstrip("/")
FOUR = {"обнаружено", "не обнаружено", "недостаточно данных", "исследовательская оценка"}
STATUS_CODES = {"detected", "not_detected", "insufficient_data", "research_estimate"}
QTY = "Количество предметов по этому снимку не определено"
BG = {"ship": r"суд|кильват", "foam": r"пен", "glint": r"блик", "cloud": r"облак", "coast": r"берег|прибо",
      "shallow": r"мелковод|мутн", "wind": r"ветер", "seam": r"шов|шв"}


def get(p):
    return urllib.request.urlopen(B + p, timeout=180).read().decode("utf-8-sig")


res = {"base": B, "critical": [], "important": [], "notes": []}
C, I, N = res["critical"].append, res["important"].append, res["notes"].append


def wind_of(p):
    for v in (p.get("wind10m_ms"), (p.get("scene") or {}).get("wind10m_ms"),
              (((p.get("probable") or {}).get("signs") or {}).get("foam") or {}).get("wind10m_ms")):
        if v is not None:
            return v
    return None


# ---------------- API
sz = json.loads(get("/api/v3/scene_zones?limit=5000"))
n_no_label = 0
for f in sz["features"]:
    p = f["properties"]
    z = p["zone_id"]
    sl = p.get("status_label")
    if sl is None:
        n_no_label += 1
    elif sl not in FOUR:
        C(f"S1 API {z}: status_label «{sl}»")
    for k in ("status_label", "detection_label"):
        if "требует проверки" in str(p.get(k) or ""):
            C(f"S1 API {z}: «требует проверки» in {k}")
    if p.get("is_find"):
        conf = p.get("confirmation")
        conf_label = p.get("confirmation_label")
        want = "совпадает с разметкой Cózar 2024 (B)" if p.get("verification") == "level_B_cozar" else "независимой разметки нет"
        # training_scene zones are a distinct legitimate case (not a mismatch): "не независимая проверка; находкой не считается"
        if conf != "training_scene" and conf_label != want:
            I(f"S2 API {z}: confirmation_label «{conf_label}» != «{want}» (confirmation={conf!r})")
    cls = p.get("classification") or {}
    exc = cls.get("excluded_backgrounds") if cls else p.get("excluded_backgrounds")
    nc = cls.get("not_checked_backgrounds") if cls else p.get("not_checked_backgrounds")
    if exc is not None:
        signs = (p.get("probable") or {}).get("signs") or {}
        flags = set(p.get("flags") or [])
        for key in exc:  # real short codes: ship/seam/foam/glint/cloud/coast/shallow/wind
            if key not in BG:
                continue
            if key == "wind":
                # NB: "excluded: wind" here is Cózar 2024's *survey-area* exclusion rule (>5 m/s -> area
                # removed from a0, debris mixed/not visible), a different rule from the foam sign's own
                # 7 m/s threshold (probable.signs.foam.rule). w>=5 here is therefore *expected*, not a bug,
                # as long as wind_note documents the distinction (else it reads as "foam ruled out").
                w = wind_of(p)
                note = str(p.get("wind_note") or "")
                if w is None:
                    C(f"S3 API {z}: «ветер» excluded but wind unknown")
                elif w >= 5 and not re.search(r"5\s*м|Cózar|Cozar|a0", note):
                    I(f"S3 API {z}: «ветер» excluded, wind {w} m/s, but wind_note doesn't explain the Cózar "
                      f"area-exclusion rule (could read as \"foam checked and ruled out\")")
                continue
            s = signs.get(key)
            if key in flags or (isinstance(s, dict) and s.get("flag")):
                C(f"S3 API {z}: «{key}» excluded but flag raised")
            elif not isinstance(s, dict):
                N(f"S3 API {z}: «{key}» excluded but no matching check in probable.signs (n/a, not necessarily a bug)")
        overlap = set(exc) & set(nc or [])
        if overlap:
            C(f"S3/B22 API {z}: same background(s) both excluded and not_checked: {sorted(overlap)}")
        if p.get("is_find"):
            nc_text = json.dumps(nc, ensure_ascii=False) + " " + str(cls.get("not_checked_label") or "")
            if not re.search(r"algae|sargassum|водоросл|саргасс", nc_text, re.I):
                I(f"S3 API {z}: algae/sargassum not in not_checked_backgrounds/_label")
    elif p.get("is_find"):
        N(f"{z}: no excluded_backgrounds yet")
    blob = json.dumps(p.get("research_estimate") or {}, ensure_ascii=False)
    if "нижняя граница" in blob:
        I(f"S4 API {z}: «нижняя граница» in research_estimate")
    if p.get("is_find"):
        e = p.get("research_estimate") or {}
        if not e.get("scenario_label") and not e.get("scenario_status"):
            N(f"{z}: no scenario_label/scenario_status yet")
if n_no_label:
    N(f"{n_no_label} zones without status_label")

# ---------------- export
rows = list(csv.DictReader(io.StringIO(get("/api/v3/export?layer=scene_zones&format=csv"))))
cols = list(rows[0].keys()) if rows else []
res["csv_cols"] = cols
if "status" in cols:
    vals = {r["status"] for r in rows if r["status"]}
    res["csv_status_values"] = sorted(vals)
    if not vals <= STATUS_CODES:
        C(f"S1 CSV status values outside the four codes: {sorted(vals - STATUS_CODES)}")
else:
    I("S5 CSV: no status column")
if "status_label" in cols:
    lvals = {r["status_label"] for r in rows if r["status_label"]}
    res["csv_status_label_values"] = sorted(lvals)
    if not lvals <= FOUR:
        C(f"S1 CSV status_label values outside the four: {sorted(lvals - FOUR)}")
else:
    I("S5 CSV: no status_label column")
for need in ("confirmation", "class", "excluded_backgrounds", "scenario_status"):
    if not any(need in c for c in cols):
        I(f"S5 CSV: no column {need}")
csvtxt = "\n".join(",".join(r.values()) for r in rows)
if "требует проверки" in csvtxt:
    I("S1 CSV: «требует проверки» appears in values (check column)")
if "нижняя граница" in csvtxt:
    I("S4 CSV: «нижняя граница» in values")
gj = get("/api/v3/export?layer=scene_zones&format=geojson")
if "нижняя граница" in gj:
    I("S4 GeoJSON: «нижняя граница»")
if re.search(r'"(status_label|detection_label)": "[^"]*требует проверки', gj):
    C("S1 GeoJSON: «требует проверки» as status")

# ---------------- docs + deck
for fn in ("README.md", "docs/SPEECH.md", "docs/DEMO.md", "docs/QA.md", "docs/QUANTITY.md", "docs/PIPELINE.md",
           "docs/CONTRACTS_V3.md", "reports/report.md"):
    fp = ROOT / fn
    if not fp.exists():
        continue
    for i, l in enumerate(fp.read_text(encoding="utf-8").split("\n"), 1):
        if "нижняя граница" in l and not re.search(r"убра|не пис|без «нижн|отмен", l):
            I(f"S4 doc {fn}:{i}: «нижняя граница»")
        if re.search(r"(статус\w*|status)[^.\n]{0,40}требует проверки", l, re.I) and not re.search(r"убра|вместо|не статус|отмен", l):
            I(f"S1 doc {fn}:{i}: «требует проверки» as status")
try:
    from pptx import Presentation
    pr = Presentation(str(ROOT / "reports/case_deck.pptx"))
    for si, s in enumerate(pr.slides, 1):
        t = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        if s.has_notes_slide:
            t.append(s.notes_slide.notes_text_frame.text)
        for x in t:
            if "требует проверки" in x:
                I(f"S1 deck slide {si}: «требует проверки»")
            if "нижняя граница" in x:
                I(f"S4 deck slide {si}: «нижняя граница»")
except Exception as ex:  # noqa: BLE001
    N(f"deck: {ex!r}"[:200])

# ---------------- UI
with sync_playwright() as pw:
    b = pw.chromium.launch(args=["--use-angle=swiftshader", "--enable-unsafe-swiftshader"])
    page = b.new_context(viewport={"width": 1920, "height": 1080}, locale="ru-RU").new_page()
    errs = []
    page.on("console", lambda m: errs.append(m.text[:200]) if m.type == "error" else None)
    page.goto(B + "/", wait_until="domcontentloaded")
    page.wait_for_function("() => window.__app && window.__app.ready && window.__app.szReady", timeout=120000)
    time.sleep(3)
    body = page.inner_text("body")
    res["start_has_trebuet"] = [m.group(0) for m in re.finditer(r".{0,60}требует проверки.{0,40}", body)][:5]
    if res["start_has_trebuet"]:
        I("S1 UI start screen / legend: «требует проверки» visible — see start_has_trebuet")
    finds = [f["properties"] for f in sz["features"] if f["properties"].get("is_find")]
    pick = [next(p for p in finds if p.get("verification") == "level_B_cozar"),
            next(p for p in finds if p.get("verification") != "level_B_cozar")]
    ins = next((f["properties"] for f in sz["features"] if f["properties"]["status"] == "insufficient_data"), None)
    if ins:
        pick.append(ins)
    res["cards"] = {}
    for p in pick:
        z = p["zone_id"]
        # fresh page per card: an expanded <details> must not carry over from the previous card
        page.goto(B + "/", wait_until="domcontentloaded")
        page.wait_for_function("() => window.__app && window.__app.ready && window.__app.szReady", timeout=120000)
        time.sleep(2)
        page.evaluate("(id) => window.__app.selectZone(id)", z)
        page.wait_for_function("() => window.__app.szDetailReady && document.querySelector('[data-testid=scene-zone-card]')",
                               timeout=60000)
        time.sleep(2.5)
        card = page.locator("[data-testid=scene-zone-card]")
        t = card.inner_text()
        page.screenshot(path=str(OUT / f"ui_s39_{z[3:]}.png"))
        lines = [l.strip() for l in t.split("\n") if l.strip()]
        qi = next((i for i, l in enumerate(lines) if l.startswith("Количество")), None)
        qty_block = " ".join(lines[qi:qi + 2]) if qi is not None else ""
        chip = page.locator("[data-testid=sz-status]").first.inner_text() if page.locator("[data-testid=sz-status]").count() else None
        vis_items = [m.group(0) for m in re.finditer(r"[≈≥~]*\s*\d[\d  ,.]*\s*шт\.?/км²", t)]
        details = page.evaluate("() => [...document.querySelectorAll('[data-testid=scene-zone-card] details')]"
                                ".map(d => ({open: d.open, summary: (d.querySelector('summary') || {}).innerText || ''}))")
        r = {"status_chip": chip, "qty_block": qty_block, "visible_items_km2": vis_items, "details": details,
             "has_confirmation": "Подтверждение" in t, "has_what": "Что это" in t, "has_excluded": "Исключено" in t,
             "has_not_checked": bool(re.search(r"Не проверяется|не проверя", t)), "lower_bound": "нижняя граница" in t}
        if chip:
            c0 = chip.strip()
            if "требует проверки" in c0:
                C(f"S1 UI {z}: «требует проверки» in status chip «{c0}»")
            elif c0 not in FOUR:
                I(f"S1 UI {z}: status chip «{c0}» is not exactly one of the four")
        if p.get("is_find"):
            if QTY not in qty_block:
                C(f"S4 UI {z}: first quantity line is «{qty_block[:120]}»")
            if vis_items:
                C(f"S4 UI {z}: items/km2 visible without expanding: {vis_items[:3]}")
            if not (r["has_what"] and r["has_excluded"]):
                I(f"S3 UI {z}: «Что это»/«Исключено» missing")
            if not r["has_not_checked"]:
                I(f"S3 UI {z}: «Не проверяется» missing")
            if not r["has_confirmation"]:
                I(f"S2 UI {z}: «Подтверждение» missing")
        elif vis_items:
            C(f"S4 UI {z} (not a find): items/km2 shown: {vis_items[:3]}")
        if r["lower_bound"]:
            I(f"S4 UI {z}: «нижняя граница»")
        for sm in page.locator("[data-testid=scene-zone-card] details summary").all():
            try:
                if re.search("сценари", sm.inner_text(), re.I):
                    sm.click()
                    time.sleep(0.5)
            except Exception:  # noqa: BLE001
                pass
        r["expanded_scenario"] = [m.group(0) for m in re.finditer(r".{0,40}шт\.?/км².{0,220}", card.inner_text())][:3]
        page.screenshot(path=str(OUT / f"ui_s39_{z[3:]}_open.png"))
        res["cards"][z] = r
    # zone list: the scene list (start) and the numbered zone list of the first scene
    page.goto(B + "/", wait_until="domcontentloaded")
    page.wait_for_function("() => window.__app && window.__app.ready && window.__app.szReady", timeout=120000)
    time.sleep(2)
    lt = page.locator("[data-testid=scene-list]").inner_text() if page.locator("[data-testid=scene-list]").count() else ""
    if page.locator("[data-testid=scene-item]").count():
        page.locator("[data-testid=scene-item]").first.click()
        time.sleep(5)
        page.screenshot(path=str(OUT / "ui_s39_scene_zone_list.png"))
        for sel in ("[data-testid=scene-zones]", "[data-testid=sz-item]"):
            if page.locator(sel).count():
                lt += " || " + " | ".join(page.locator(sel).nth(i).inner_text() for i in range(min(page.locator(sel).count(), 30)))
                break
    res["list_sample"] = lt[:800]
    if re.search(r"шт\.?/км²", lt):
        C("S4 UI list: items/km2 shown in the zone list")
    if "требует проверки" in lt:
        C("S1 UI list: «требует проверки» in the zone list")
    res["console_errors"] = errs[:10]
    b.close()

res["critical"] = sorted(set(res["critical"]))
res["important"] = sorted(set(res["important"]))
(OUT / "s39.json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
short = {k: (v[:12] + [f"... {len(v)} total"] if isinstance(v, list) and len(v) > 12 else v)
         for k, v in res.items() if k != "csv_cols"}
print(json.dumps(short, ensure_ascii=False, indent=1)[:7000])
